"""Excel-Liste aller Arbeitsplaetze zur Sektorpflege durch die Fertigungssteuerung.

Statt ~10.000 Vorgangsuebergaengen (alte Matrix-Vorlage) muss nur noch jeder
Arbeitsplatz einmal einem Sektor zugeordnet werden.
"""

from datetime import date

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text

from ..config import PROJECT_ROOT

SQL = text("""
    SELECT a.work_cntr AS Arbeitsplatz, a.bezeichnung AS Bezeichnung, a.res_typ AS Typ,
           a.sektor AS Sektor, a.transport_modus AS Transport_Modus, a.lagerort_code AS Lagerort
    FROM sce_mes.arbeitsplatz a
    ORDER BY a.bezeichnung
""")


def export(engine) -> str:
    with engine.connect() as con:
        df = pd.read_sql(SQL, con)
    out = PROJECT_ROOT / "output" / f"arbeitsplaetze_sektoren_{date.today():%Y%m%d}.xlsx"
    out.parent.mkdir(exist_ok=True)
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="Arbeitsplaetze")
        ws = xw.sheets["Arbeitsplaetze"]
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="003C7D")
        gelb = PatternFill("solid", fgColor="FFF9C4")
        for row in range(2, ws.max_row + 1):
            for col in (4, 5, 6):
                ws.cell(row=row, column=col).fill = gelb
        dv = DataValidation(type="list", formula1='"teil,voll"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"E2:E{max(ws.max_row, 2)}")
        for col, w in zip("ABCDEF", (14, 40, 12, 16, 16, 16)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
    return str(out)
